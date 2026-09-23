"""The per-release mapping report: the columns Lakshya's requirement asks for.

The report is a view of the same Resolution objects the canonical mapping is
written from. These tests hold that line: the twenty required columns are
present and in order, an unknown field stays empty rather than being filled in
with something plausible, and no row that is not VERIFIED goes out without a
reason attached.
"""

from __future__ import annotations

import csv

from openpyxl import load_workbook

from apm.models import (
    Category, Confidence, DebianSource, Method, Resolution, Status, Upstream,
)
from apm.report import RELEASE_COLUMNS, release_row, write_release_reports

REQUIRED_IN_ORDER = [
    "Package", "ARCoS Repository", "ARCoS Path", "ARCoS Release", "ARCoS Branch",
    "Debian Release", "Debian Source Package", "Debian Version",
    "Debian VCS Repository", "Debian VCS Branch", "Upstream Repository",
    "Upstream Branch", "Upstream Tag", "Resolution Mode", "Resolution Status",
    "Resolution Reason", "Evidence Source", "Evidence URL",
    "Verification Method", "Common Ancestor",
]


def _verified(**overrides) -> Resolution:
    resolution = Resolution(
        package="iputils", release="bookworm", category=Category.DEBIAN_UPSTREAM,
        status=Status.VERIFIED, method=Method.DEP12, confidence=Confidence.HIGH,
        arcos_repository="ssh://git@github.com/Arrcus/iputils.git",
        github_repository="Arrcus/iputils", arcos_branch="aminor",
        arcos_path="packages/iputils", arcos_release="aminor",
        arcos_commit="9" * 40,
        debian=DebianSource(
            package="iputils", version="3:20221126-1+deb12u1", directory="pool/x",
            suite="bookworm",
            vcs_git="https://salsa.debian.org/debian/iputils.git",
            vcs_branch="debian/master",
        ),
        upstream=Upstream(repository="https://github.com/iputils/iputils.git",
                          ref="master"),
        resolved_sha="1" * 40, merge_base="a" * 40, origin_kind="project",
        evidence_source="debian/upstream/metadata (DEP-12) Repository:",
        evidence_url="https://github.com/iputils/iputils.git",
        verification="git merge-base: the fork shares history with master",
    )
    for key, value in overrides.items():
        setattr(resolution, key, value)
    return resolution


def test_the_required_columns_come_first_and_in_order():
    assert RELEASE_COLUMNS[:len(REQUIRED_IN_ORDER)] == REQUIRED_IN_ORDER


def test_a_verified_row_carries_repository_branch_and_evidence():
    row = release_row(_verified())
    assert row["Upstream Repository"] == "https://github.com/iputils/iputils.git"
    assert row["Upstream Branch"] == "master"
    assert row["Upstream Tag"] == ""
    assert row["Evidence Source"].startswith("debian/upstream/metadata")
    assert row["Verification Method"].startswith("git merge-base")
    assert row["Common Ancestor"] == "a" * 12
    assert row["ARCoS Path"] == "packages/iputils"
    assert row["ARCoS Release"] == "aminor"


def test_a_tag_is_reported_as_a_tag_and_not_as_a_branch():
    resolution = _verified(
        upstream=Upstream(repository="https://example.invalid/x.git",
                          ref="v1.2.3", is_tag=True)
    )
    row = release_row(resolution)
    assert row["Upstream Tag"] == "v1.2.3"
    assert row["Upstream Branch"] == ""


def test_the_debian_packaging_repository_is_never_the_upstream_column():
    row = release_row(_verified())
    assert row["Debian VCS Repository"] == "https://salsa.debian.org/debian/iputils.git"
    assert row["Debian VCS Branch"] == "debian/master"
    assert "salsa" not in row["Upstream Repository"]


def test_an_unknown_field_stays_empty_rather_than_being_invented():
    resolution = _verified(debian=None, upstream=None, merge_base=None,
                           evidence_source="", evidence_url="", verification="")
    row = release_row(resolution)
    for column in ("Debian Source Package", "Debian Version",
                   "Debian VCS Repository", "Debian VCS Branch",
                   "Upstream Repository", "Upstream Branch", "Upstream Tag",
                   "Common Ancestor", "Evidence Source", "Evidence URL",
                   "Verification Method"):
        assert row[column] == "", column


def test_every_unverified_row_explains_itself():
    resolution = _verified(status=Status.NEEDS_REVIEW)
    resolution.note("The fork shares no history with any candidate upstream.")
    assert release_row(resolution)["Resolution Reason"].startswith("The fork shares")


def test_a_no_upstream_row_is_an_answer_not_an_error():
    resolution = _verified(
        status=Status.NO_UPSTREAM, category=Category.ARRCUS_NATIVE,
        upstream=None, method=Method.NOT_APPLICABLE,
    )
    resolution.note("No upstream exists for this package; it is arrcus native.")
    row = release_row(resolution)
    assert row["Resolution Status"] == "NO_UPSTREAM"
    assert row["Resolution Reason"]
    assert row["Upstream Repository"] == ""


def test_only_the_named_release_is_written(tmp_path):
    rows = [_verified(), _verified(release="trixie")]
    written = write_release_reports(rows, tmp_path, "bookworm")

    assert written["rows"] == 1
    with (tmp_path / "upstream-mapping-bookworm.csv").open() as handle:
        records = list(csv.DictReader(handle))
    assert [r["Debian Release"] for r in records] == ["bookworm"]


def test_the_workbook_has_mapping_summary_and_needs_review(tmp_path):
    needs_review = _verified(package="babeltrace", status=Status.NEEDS_REVIEW)
    needs_review.note("No candidate shared history with the fork.")
    write_release_reports([_verified(), needs_review], tmp_path, "bookworm")

    book = load_workbook(tmp_path / "upstream-mapping-bookworm.xlsx")
    assert book.sheetnames == ["Mapping", "Summary", "Needs Review"]

    review = book["Needs Review"]
    packages = [row[0] for row in review.iter_rows(min_row=2, values_only=True)]
    assert packages == ["babeltrace"], "verified rows do not belong in the review sheet"
