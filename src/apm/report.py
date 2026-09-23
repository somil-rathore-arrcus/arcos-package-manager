"""Write the mapping out as CSV and as a reviewable workbook.

The CSV is the diffable record - it is what a later run is compared against. The
workbook is for reading: three sheets, colour keyed to status, and every
identifier hyperlinked to the page that backs it.
"""

from __future__ import annotations

import csv
import logging
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import links
from .models import Resolution, Status

log = logging.getLogger(__name__)

COLUMNS = [
    "Package", "Category", "Debian Release", "Status", "Confidence",
    "Resolution Mode", "Resolution Method",
    "ARCoS Repository", "ARCoS Path", "ARCoS Release", "ARCoS Branch",
    "ARCoS Commit",
    "Debian Source Package", "Debian Version",
    "Upstream Repository", "Upstream Ref", "Upstream Branch", "Upstream Tag",
    "Origin Kind", "Upstream Commit",
    "Commits Behind", "ARCoS-only Commits", "Merge Base",
    "Vcs-Git (packaging)", "Vcs-Git Branch", "Homepage",
    "Evidence Source", "Evidence URL", "Verification", "Reason",
    "ARCoS Link", "Debian Link", "Upstream Link",
    "Evidence", "Notes",
]

# The columns Lakshya's Bookworm/aminor requirement asks for, in that order.
# A view of the same Resolution objects the canonical mapping is written from -
# not a second mapping, which would be free to disagree with the first.
RELEASE_COLUMNS = [
    "Package",
    "ARCoS Repository",
    "ARCoS Path",
    "ARCoS Release",
    "ARCoS Branch",
    "Debian Release",
    "Debian Source Package",
    "Debian Version",
    "Debian VCS Repository",
    "Debian VCS Branch",
    "Upstream Repository",
    "Upstream Branch",
    "Upstream Tag",
    "Resolution Mode",
    "Resolution Status",
    "Resolution Reason",
    "Evidence Source",
    "Evidence URL",
    "Verification Method",
    "Common Ancestor",
    # Beyond the required twenty, and worth keeping: what kind of package it is,
    # which side of the fork the upstream is, and the full evidence trail.
    "Category",
    "Confidence",
    "Resolution Method",
    "Origin Kind",
    "ARCoS Commit",
    "Upstream Commit",
    "Commits Behind",
    "ARCoS-only Commits",
    "ARCoS Link",
    "Debian Link",
    "Upstream Link",
    "Evidence",
    "Notes",
]

# Columns rendered as clickable links, by header name.
LINK_COLUMNS = {"ARCoS Link", "Debian Link", "Upstream Link"}

STATUS_FILLS = {
    Status.VERIFIED.value: "D9EAD3",      # green
    Status.NEEDS_REVIEW.value: "FFF2CC",  # amber
    Status.UNRESOLVED.value: "F4CCCC",    # red
    Status.NO_UPSTREAM.value: "EFEFEF",   # grey - a finished answer, not a problem
}

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)
LINK_FONT = Font(color="0563C1", underline="single")


def to_row(resolution: Resolution) -> dict:
    debian = resolution.debian
    upstream = resolution.upstream
    return {
        "Package": resolution.package,
        "Category": resolution.category.value,
        "Debian Release": resolution.release,
        "Status": resolution.status.value,
        "Confidence": resolution.confidence.value,
        "Resolution Mode": resolution.mode,
        "Resolution Method": resolution.method.value,
        "ARCoS Repository": resolution.github_repository,
        "ARCoS Path": resolution.arcos_path,
        "ARCoS Release": resolution.arcos_release,
        "ARCoS Branch": resolution.arcos_branch,
        "ARCoS Commit": (resolution.arcos_commit or "")[:12],
        "Debian Source Package": debian.package if debian else "",
        "Debian Version": debian.version if debian else "",
        "Upstream Repository": upstream.repository if upstream else "",
        "Upstream Ref": (upstream.ref or "") if upstream else "",
        "Upstream Branch": resolution.upstream_branch,
        "Upstream Tag": resolution.upstream_tag,
        "Origin Kind": resolution.origin_kind or "",
        "Upstream Commit": (resolution.resolved_sha or "")[:12],
        "Commits Behind": "" if resolution.behind is None else resolution.behind,
        "ARCoS-only Commits": (
            "" if resolution.arcos_only is None else resolution.arcos_only
        ),
        "Merge Base": (resolution.merge_base or "")[:12],
        "Vcs-Git (packaging)": (debian.vcs_git or "") if debian else "",
        "Vcs-Git Branch": (debian.vcs_branch or "") if debian else "",
        "Homepage": (debian.homepage or "") if debian else "",
        "Evidence Source": resolution.evidence_source,
        "Evidence URL": resolution.evidence_url,
        "Verification": resolution.verification,
        "Reason": resolution.reason,
        "ARCoS Link": links.arcos_commit_web(
            resolution.github_repository, resolution.arcos_commit
        ) or links.arcos_web(resolution.github_repository, resolution.arcos_branch) or "",
        "Debian Link": links.debian_source_web(
            resolution.release, debian.package if debian else None
        ) or "",
        "Upstream Link": links.upstream_web(
            upstream.repository if upstream else None,
            upstream.ref if upstream else None,
        ) or "",
        "Evidence": "\n".join(resolution.evidence),
        "Notes": "\n".join(resolution.notes),
    }


def release_row(resolution: Resolution) -> dict:
    """The requirement's view of one resolution.

    Built from the same Resolution object `to_row` uses, so the two reports
    cannot drift: there is one mapping, rendered twice.
    """
    row = to_row(resolution)
    debian = resolution.debian
    return {
        "Package": row["Package"],
        "ARCoS Repository": row["ARCoS Repository"],
        "ARCoS Path": row["ARCoS Path"],
        "ARCoS Release": row["ARCoS Release"],
        "ARCoS Branch": row["ARCoS Branch"],
        "Debian Release": row["Debian Release"],
        "Debian Source Package": row["Debian Source Package"],
        "Debian Version": row["Debian Version"],
        "Debian VCS Repository": (debian.vcs_git or "") if debian else "",
        "Debian VCS Branch": (debian.vcs_branch or "") if debian else "",
        "Upstream Repository": row["Upstream Repository"],
        "Upstream Branch": row["Upstream Branch"],
        "Upstream Tag": row["Upstream Tag"],
        "Resolution Mode": row["Resolution Mode"],
        "Resolution Status": row["Status"],
        "Resolution Reason": row["Reason"],
        "Evidence Source": row["Evidence Source"],
        "Evidence URL": row["Evidence URL"],
        "Verification Method": row["Verification"],
        "Common Ancestor": row["Merge Base"],
        "Category": row["Category"],
        "Confidence": row["Confidence"],
        "Resolution Method": row["Resolution Method"],
        "Origin Kind": row["Origin Kind"],
        "ARCoS Commit": row["ARCoS Commit"],
        "Upstream Commit": row["Upstream Commit"],
        "Commits Behind": row["Commits Behind"],
        "ARCoS-only Commits": row["ARCoS-only Commits"],
        "ARCoS Link": row["ARCoS Link"],
        "Debian Link": row["Debian Link"],
        "Upstream Link": row["Upstream Link"],
        "Evidence": row["Evidence"],
        "Notes": row["Notes"],
    }


def write_release_reports(resolutions: list, out_dir: Path, release: str) -> dict:
    """Write out/upstream-mapping-<release>.{csv,xlsx} for one release.

    Only that release's rows: a package absent from the release manifest has no
    row here at all, because it does not ship in that release.
    """
    rows = [release_row(r) for r in resolutions if r.release == release]
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    csv_path = out_dir / f"upstream-mapping-{release}.csv"
    xlsx_path = out_dir / f"upstream-mapping-{release}.xlsx"

    with csv_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=RELEASE_COLUMNS)
        writer.writeheader()
        writer.writerows(rows)

    book = Workbook()
    _release_mapping_sheet(book.active, rows)
    _release_summary_sheet(
        book.create_sheet("Summary"),
        [r for r in resolutions if r.release == release],
        release,
    )
    _release_review_sheet(book.create_sheet("Needs Review"), rows)
    book.save(xlsx_path)

    log.info("wrote %s and %s (%d rows)", csv_path, xlsx_path, len(rows))
    return {"csv": csv_path, "xlsx": xlsx_path, "rows": len(rows)}


_RELEASE_LINK_COLUMNS = {
    "ARCoS Repository", "ARCoS Branch", "Debian VCS Repository",
    "Debian VCS Branch", "Upstream Repository", "Upstream Branch",
    "Upstream Tag", "Evidence URL", "ARCoS Link", "Debian Link",
    "Upstream Link",
}


def _release_link(header: str, row: dict) -> str:
    """The URL a cell should point at, or '' when it cannot be justified."""
    value = row.get(header) or ""
    if header == "ARCoS Repository":
        return links.arcos_web(value, row.get("ARCoS Branch")) or ""
    if header in ("Debian VCS Repository", "Upstream Repository"):
        return links.upstream_web(value, None) or "" if value else ""
    # Branch and tag cells point at that branch, not just at the repository:
    # "which branch" is the question this sheet exists to answer.
    if header == "ARCoS Branch":
        return links.arcos_web(row.get("ARCoS Repository", ""), value) or ""
    if header in ("Upstream Branch", "Upstream Tag"):
        return (
            links.upstream_web(row.get("Upstream Repository", ""), value) or ""
        ) if value else ""
    if header == "Debian VCS Branch":
        return (
            links.upstream_web(row.get("Debian VCS Repository", ""), value) or ""
        ) if value else ""
    return value if value.startswith("http") else ""


def _release_mapping_sheet(sheet, rows: list) -> None:
    sheet.title = "Mapping"
    _write_header(sheet, RELEASE_COLUMNS)
    status_index = RELEASE_COLUMNS.index("Resolution Status") + 1
    for row in rows:
        sheet.append([row.get(c, "") for c in RELEASE_COLUMNS])
        excel_row = sheet.max_row
        fill = STATUS_FILLS.get(row.get("Resolution Status", ""))
        if fill:
            sheet.cell(row=excel_row, column=status_index).fill = PatternFill(
                "solid", fgColor=fill
            )
        for header in _RELEASE_LINK_COLUMNS:
            target = _release_link(header, row)
            if not target:
                continue
            cell = sheet.cell(
                row=excel_row, column=RELEASE_COLUMNS.index(header) + 1
            )
            cell.hyperlink = target
            cell.font = LINK_FONT
    sheet.auto_filter.ref = sheet.dimensions
    _fit_columns(
        sheet, RELEASE_COLUMNS,
        wide={"Evidence": 60, "Notes": 45, "Resolution Reason": 50,
              "Verification Method": 50, "Evidence Source": 34},
    )


def _release_summary_sheet(sheet, resolutions: list, release: str) -> None:
    from collections import Counter

    sheet.append([f"ARCoS upstream mapping - Debian {release}"])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append([])

    arcos_releases = sorted({r.arcos_release for r in resolutions if r.arcos_release})
    sheet.append(["ARCoS release (manifest branch)", ", ".join(arcos_releases)])
    sheet.append(["Packages in this release", len(resolutions)])
    sheet.append([])

    for title, values in (
        ("Status", [r.status.value for r in resolutions]),
        ("Category", [r.category.value for r in resolutions]),
        ("Resolution method", [r.method.value for r in resolutions]),
        ("Origin kind", [r.origin_kind for r in resolutions if r.origin_kind]),
    ):
        sheet.append([title, "Packages"])
        for cell in sheet[sheet.max_row]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
        for value, count in sorted(Counter(values).items()):
            sheet.append([value, count])
            if title == "Status":
                colour = STATUS_FILLS.get(value)
                if colour:
                    sheet.cell(row=sheet.max_row, column=1).fill = PatternFill(
                        "solid", fgColor=colour
                    )
        sheet.append([])

    for line in (
        "VERIFIED: the upstream repository was contacted, its ref resolved to a "
        "commit, and where a fork was probed the shared history was proven.",
        "NEEDS_REVIEW: an answer exists but is not proven - the reason column "
        "says what is missing.",
        "NO_UPSTREAM: searched, and there is provably nothing to track. A "
        "finished answer, not a failure.",
        "Upstream branch and upstream tag are separate columns: a tag is a fixed "
        "point, a branch keeps moving, and they are not interchangeable.",
    ):
        sheet.append([line])
    sheet.column_dimensions["A"].width = 42
    sheet.column_dimensions["B"].width = 18


def _release_review_sheet(sheet, rows: list) -> None:
    columns = [
        "Package", "ARCoS Repository", "ARCoS Branch", "Debian Source Package",
        "Debian Version", "Debian VCS Repository", "Upstream Repository",
        "Upstream Branch", "Resolution Status", "Resolution Reason",
        "Evidence Source", "Evidence URL", "Evidence",
    ]
    _write_header(sheet, columns)
    for row in rows:
        if row.get("Resolution Status") == Status.VERIFIED.value:
            continue
        sheet.append([row.get(c, "") for c in columns])
    sheet.auto_filter.ref = sheet.dimensions
    _fit_columns(
        sheet, columns,
        wide={"Evidence": 70, "Resolution Reason": 60, "Evidence Source": 30},
    )


def write_csv(resolutions: list, path: Path) -> None:
    write_csv_rows([to_row(r) for r in resolutions], path)


def write_xlsx(resolutions: list, path: Path) -> None:
    write_xlsx_rows([to_row(r) for r in resolutions], path)


def write_csv_rows(rows: list, path: Path) -> None:
    """Write already-rendered rows.

    Taking rows rather than Resolution objects is what lets a single-release run
    keep the other releases' rows from the existing mapping: they are carried
    through exactly as they were written, with nothing re-derived from them.
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=COLUMNS)
        writer.writeheader()
        for row in rows:
            writer.writerow({column: row.get(column, "") for column in COLUMNS})
    log.info("wrote %s (%d rows)", path, len(rows))


def write_xlsx_rows(rows: list, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    _mapping_sheet(book.active, rows)
    _summary_sheet(book.create_sheet("Summary"), rows)
    _review_sheet(book.create_sheet("Needs review"), rows)
    book.save(path)
    log.info("wrote %s (%d rows)", path, len(rows))


def _write_header(sheet, columns: list) -> None:
    sheet.append(columns)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"


def _mapping_sheet(sheet, rows: list) -> None:
    sheet.title = "Mapping"
    _write_header(sheet, COLUMNS)

    status_index = COLUMNS.index("Status") + 1
    for row in rows:
        # .get: rows carried over from an older mapping predate newer columns.
        sheet.append([row.get(c, "") for c in COLUMNS])
        excel_row = sheet.max_row
        fill_colour = STATUS_FILLS.get(row.get("Status", ""))
        if fill_colour:
            sheet.cell(row=excel_row, column=status_index).fill = PatternFill(
                "solid", fgColor=fill_colour
            )
        for header in LINK_COLUMNS:
            value = row[header]
            if not value:
                continue
            cell = sheet.cell(row=excel_row, column=COLUMNS.index(header) + 1)
            cell.hyperlink = value
            cell.font = LINK_FONT

    sheet.auto_filter.ref = sheet.dimensions
    _fit_columns(sheet, COLUMNS, wide={"Evidence": 60, "Notes": 45})


def _summary_sheet(sheet, rows: list) -> None:
    from collections import Counter

    releases = sorted({r.get("Debian Release", "") for r in rows})

    sheet.append(["ARCoS package upstream coverage"])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append([])

    def block(title: str, key: str, order=None) -> None:
        sheet.append([title] + releases + ["Total"])
        for cell in sheet[sheet.max_row]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
        counts = Counter(
            (r.get(key, ""), r.get("Debian Release", "")) for r in rows if r.get(key)
        )
        names = order or sorted({r.get(key, "") for r in rows if r.get(key)})
        for name in names:
            values = [counts.get((name, rel), 0) for rel in releases]
            if not any(values):
                continue
            sheet.append([name] + values + [sum(values)])
            if key == "Status":
                fill = STATUS_FILLS.get(name)
                if fill:
                    sheet.cell(row=sheet.max_row, column=1).fill = PatternFill(
                        "solid", fgColor=fill
                    )
        sheet.append([])

    block("Status", "Status", [s.value for s in Status])
    sheet.append(["Total"] + [
        sum(1 for r in rows if r.get("Debian Release") == rel) for rel in releases
    ] + [len(rows)])
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
    sheet.append([])

    block("Category", "Category")
    block("Resolution method", "Resolution Method")
    block("Origin kind", "Origin Kind")

    sheet.append([
        "Origin kind 'debian_packaging' means the ARCoS fork descends from the "
        "Debian packaging repository, so its patches are packaging changes."
    ])
    sheet.append([
        "NO_UPSTREAM is a finished answer: the package is Arrcus-native, a vendor "
        "drop, or was searched for and has no public upstream."
    ])
    sheet.append([
        "NEEDS_REVIEW means an upstream was found and verified to exist, but the "
        "branch was inferred rather than confirmed against the fork's history."
    ])
    for column, width in (("A", 26), ("B", 12), ("C", 12), ("D", 10)):
        sheet.column_dimensions[column].width = width


def _review_sheet(sheet, rows: list) -> None:
    columns = [
        "Package", "Debian Release", "Status", "Confidence", "Resolution Method",
        "Upstream Repository", "Upstream Ref", "Origin Kind", "Commits Behind",
        "Upstream Link", "Evidence", "Notes",
    ]
    _write_header(sheet, columns)
    for row in rows:
        if row.get("Status") == Status.VERIFIED.value:
            continue
        if row.get("Status") == Status.NO_UPSTREAM.value:
            continue
        sheet.append([row.get(c, "") for c in columns])
        cell = sheet.cell(row=sheet.max_row, column=columns.index("Upstream Link") + 1)
        if cell.value:
            cell.hyperlink = cell.value
            cell.font = LINK_FONT
    sheet.auto_filter.ref = sheet.dimensions
    _fit_columns(sheet, columns, wide={"Evidence": 70, "Notes": 45})


def _fit_columns(sheet, columns: list, wide: dict) -> None:
    for index, header in enumerate(columns, start=1):
        letter = get_column_letter(index)
        if header in wide:
            sheet.column_dimensions[letter].width = wide[header]
            continue
        longest = max(
            [len(header)]
            + [
                len(str(sheet.cell(row=r, column=index).value or "").split("\n")[0])
                for r in range(2, min(sheet.max_row, 200) + 1)
            ]
        )
        sheet.column_dimensions[letter].width = min(max(longest + 2, 10), 46)
